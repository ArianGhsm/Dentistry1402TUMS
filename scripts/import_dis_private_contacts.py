#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, Optional, Tuple

from openpyxl import load_workbook


FA_TO_EN = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789",
)


def normalize_digits(value: object) -> str:
    return str(value or "").translate(FA_TO_EN)


def digits_only(value: object) -> str:
    return "".join(ch for ch in normalize_digits(value) if ch.isdigit())


def normalize_student_number(value: object) -> str:
    return digits_only(value)


def normalize_national_code(value: object) -> str:
    digits = digits_only(value)
    return digits if len(digits) == 10 else ""


def normalize_phone_number(value: object) -> str:
    digits = digits_only(value)
    if not digits:
        return ""

    if digits.startswith("0098"):
        digits = digits[4:]
    elif digits.startswith("98"):
        digits = digits[2:]

    digits = digits.lstrip("0")
    if len(digits) != 10 or not digits.startswith("9"):
        return ""
    return f"+98{digits}"


def header_row_and_indices(rows: Iterable[Tuple[object, ...]]) -> Tuple[int, Dict[str, int]]:
    key_student = "\u062f\u0627\u0646\u0634\u062c\u0648\u06cc\u06cc"  # دانشجویی
    key_personnel = "\u067e\u0631\u0633\u0646\u0644\u06cc"  # پرسنلی
    key_national = "\u06a9\u062f \u0645\u0644\u06cc"  # کد ملی
    key_phone = "\u0634\u0645\u0627\u0631\u0647 \u062a\u0644\u0641\u0646"  # شماره تلفن

    for idx, row in enumerate(rows, start=1):
        normalized = [normalize_digits(cell).strip() for cell in row]
        lowered = [cell.lower() for cell in normalized]
        col_student = next(
            (i for i, val in enumerate(lowered) if key_student in val or key_personnel in val),
            None,
        )
        col_national = next((i for i, val in enumerate(lowered) if key_national in val), None)
        col_phone = next((i for i, val in enumerate(lowered) if key_phone in val), None)

        if col_student is not None and col_national is not None and col_phone is not None:
            return idx, {"student": col_student, "national": col_national, "phone": col_phone}

    raise RuntimeError("Header row with student/national/phone columns was not found in workbook.")


def load_dis_rows(xlsx_path: Path) -> Dict[str, Dict[str, str]]:
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    header_row, cols = header_row_and_indices(all_rows[:40])

    out: Dict[str, Dict[str, str]] = {}
    for row in all_rows[header_row:]:
        student = normalize_student_number(row[cols["student"]] if cols["student"] < len(row) else "")
        if not student:
            continue
        national = normalize_national_code(row[cols["national"]] if cols["national"] < len(row) else "")
        phone = normalize_phone_number(row[cols["phone"]] if cols["phone"] < len(row) else "")
        out[student] = {"nationalCode": national, "directoryPhoneNumber": phone}
    return out


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def merge_private_fields(
    users_path: Path,
    rows: Dict[str, Dict[str, str]],
    overwrite_existing: bool,
) -> Dict[str, int]:
    payload = json.loads(users_path.read_text(encoding="utf-8"))
    users = payload.get("users", {})
    if not isinstance(users, dict):
        raise RuntimeError("Invalid users.json structure: 'users' object is missing.")

    stats = {"rows": len(rows), "matchedUsers": 0, "updatedUsers": 0, "notFoundUsers": 0}
    timestamp = now_iso()

    for student_number, incoming in rows.items():
        user = users.get(student_number)
        if not isinstance(user, dict):
            stats["notFoundUsers"] += 1
            continue

        stats["matchedUsers"] += 1
        changed = False

        incoming_national = incoming.get("nationalCode", "")
        current_national = str(user.get("nationalCode", "") or "")
        if incoming_national and (overwrite_existing or not current_national):
            if current_national != incoming_national:
                user["nationalCode"] = incoming_national
                changed = True

        incoming_phone = incoming.get("directoryPhoneNumber", "")
        current_phone = str(user.get("directoryPhoneNumber", "") or "")
        if incoming_phone and (overwrite_existing or not current_phone):
            if current_phone != incoming_phone:
                user["directoryPhoneNumber"] = incoming_phone
                changed = True

        if changed:
            user["updatedAt"] = timestamp
            stats["updatedUsers"] += 1

    users_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=4) + "\n",
        encoding="utf-8",
    )
    return stats


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Import national code and directory phone fields from DIS xlsx into auth users store.",
    )
    parser.add_argument("--xlsx", required=True, help="Path to DIS Excel file.")
    parser.add_argument(
        "--users",
        default="server-only/storage/auth/users.json",
        help="Path to users.json store.",
    )
    parser.add_argument(
        "--overwrite-existing",
        action="store_true",
        help="Overwrite existing nationalCode/directoryPhoneNumber if present.",
    )
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx).expanduser().resolve()
    users_path = Path(args.users).expanduser().resolve()
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Excel file not found: {xlsx_path}")
    if not users_path.exists():
        raise FileNotFoundError(f"users.json not found: {users_path}")

    rows = load_dis_rows(xlsx_path)
    stats = merge_private_fields(users_path, rows, overwrite_existing=args.overwrite_existing)
    print(
        "Import complete | rows={rows} matched={matchedUsers} updated={updatedUsers} not_found={notFoundUsers}".format(
            **stats
        )
    )


if __name__ == "__main__":
    main()
